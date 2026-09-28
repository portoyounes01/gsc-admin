from __future__ import annotations

import re
import unicodedata
from datetime import date
from io import BytesIO
from typing import Iterable

import pandas as pd
import pdfplumber
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


PATIENT_NR_COL = "Patient nr."
PATIENT_DONE_COL = "Patient behandeld"
PATIENT_COL = "Patient"
TREATMENTS_COL = "Behandeling(en)"
DATE_PERIOD_COL = "Datum of periode"
REMARKS_COL = "Opmerkingen"
CHECKED_COL = "Afgevinkt"
CHECKBOX_DONE = "x"
DUPLICATE_PATIENT_NAME_REMARK = "Dubbele naam in selectie; afzonderlijk gehouden per PDF."
SCREENING_TREATMENT = "Screening"
SCREENING_DONE_COL = "Patient gescreend"

PREVIEW_COLUMNS = [
    PATIENT_DONE_COL,
    PATIENT_NR_COL,
    PATIENT_COL,
    TREATMENTS_COL,
    DATE_PERIOD_COL,
    REMARKS_COL,
]


def clean_text(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", str(value or "").replace("\r", "\n")).strip()


def simplify_for_matching(value: str) -> str:
    return (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


def format_european_date(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def format_duration(minutes: int) -> str:
    if minutes <= 0:
        return "0 min"
    hours, remaining_minutes = divmod(minutes, 60)
    if not hours:
        return f"{remaining_minutes} min"
    if not remaining_minutes:
        return f"{hours}u"
    return f"{hours}u {remaining_minutes}min"


def patient_name_from_filename(source_name: str) -> str:
    stem = source_name.rsplit(".", 1)[0]
    return re.sub(r"[_-]+", " ", stem).strip() or "Onbekende patient"


def clean_patient_name(value: str) -> str:
    cleaned = re.sub(r"^(beste|geachte)\s+", "", str(value or "").strip(), flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", cleaned).strip()


def extract_pdf_text(file_bytes: bytes) -> tuple[str, str]:
    try:
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            text_by_page = [page.extract_text() or "" for page in pdf.pages]
        text = clean_text("\n".join(text_by_page))
        if not text:
            return "", "Geen leesbare tekst gevonden. De PDF is mogelijk gescand."
        return text, ""
    except Exception as exc:
        return "", f"PDF kon niet gelezen worden: {exc}"


def find_labeled_value(text: str, labels: Iterable[str]) -> str:
    label_pattern = "|".join(re.escape(label) for label in labels)
    pattern = re.compile(rf"(?:{label_pattern})\s*[:\-]\s*(.+)", re.IGNORECASE)
    for line in text.splitlines():
        match = pattern.search(line.strip())
        if match:
            return match.group(1).strip()
    return ""


def guess_patient_name(text: str, fallback: str) -> tuple[str, str]:
    labeled_name = find_labeled_value(
        text,
        [
            "patient",
            "patientnaam",
            "naam patient",
            "naam en voornaam",
            "bewoner",
        ],
    )
    if labeled_name:
        return clean_patient_name(labeled_name), ""

    for line in text.splitlines()[:12]:
        cleaned = line.strip()
        if 4 <= len(cleaned) <= 80 and not re.search(r"\d", cleaned):
            lowered = simplify_for_matching(cleaned)
            if "behandelplan" not in lowered and "global smile" not in lowered:
                return clean_patient_name(cleaned), "Patientnaam geschat uit documentkop."
    return fallback, "Patientnaam uit bestandsnaam gehaald."


def extract_behandelplan_treatment_lines(text: str) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    start_index = None
    for index, line in enumerate(lines):
        if simplify_for_matching(line) == "tand behandeling detail opmerking":
            start_index = index + 1
            break
    if start_index is None:
        return []

    stop_words = ("total", "totaal", "handtekening", "opmerking planning")
    result = []
    for line in lines[start_index:]:
        if any(word in simplify_for_matching(line) for word in stop_words):
            break
        result.append(line)
    return result


def collapse_full_cleaning(treatments: list[str]) -> list[str]:
    quadrant_pattern = re.compile(
        r"^Kw\s+(I|II|III|IV)\s+-\s+Reiniging(?:\s+(?:Kw\s+)?\1)?$",
        re.IGNORECASE,
    )
    quadrants = {
        match.group(1).upper()
        for treatment in treatments
        if (match := quadrant_pattern.match(treatment.strip()))
    }
    if quadrants != {"I", "II", "III", "IV"}:
        return treatments

    collapsed = []
    added_full_cleaning = False
    for treatment in treatments:
        if quadrant_pattern.match(treatment.strip()):
            if not added_full_cleaning:
                collapsed.append("Volledige Reiniging")
                added_full_cleaning = True
            continue
        collapsed.append(treatment)
    return collapsed


def normalize_treatments(lines: list[str]) -> list[str]:
    treatments = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        next_line = lines[index + 1].strip() if index + 1 < len(lines) else ""
        tooth_match = re.match(r"^(\d{1,2}(?:-\d{1,2})?|Kw\s+[IVX]+)\s+(.+)$", line, re.IGNORECASE)
        if tooth_match:
            tooth, treatment = tooth_match.groups()
            treatments.append(f"{tooth} - {treatment.strip()}")
        elif line.lower() == "cofferdam" and treatments:
            treatments[-1] = f"{treatments[-1]} - {line}"
        elif line and not re.match(r"^(tand|behandeling|detail|opmerking)$", line, re.IGNORECASE):
            if next_line and line.lower() in {"herstelling", "gebitsprothese", "prothese"}:
                treatments.append(f"{line} {next_line}")
                index += 1
            else:
                treatments.append(line)
        index += 1
    return collapse_full_cleaning(list(dict.fromkeys(treatments)))


def guess_treatments(text: str) -> tuple[list[str], str]:
    table_lines = extract_behandelplan_treatment_lines(text)
    if table_lines:
        treatments = normalize_treatments(table_lines)
        if treatments:
            return treatments, ""

    labeled = find_labeled_value(text, ["behandeling", "behandelingen", "treatment", "treatments"])
    if labeled:
        return split_display_treatments(labeled), ""
    return ["Behandeling ontbreekt"], "Behandelingstabel niet gevonden of niet herkend."


def unique_non_empty(values: Iterable[str]) -> list[str]:
    cleaned_values = []
    for value in values:
        cleaned = str(value or "").strip()
        if cleaned and cleaned.lower() != "nan":
            cleaned_values.append(cleaned)
    return list(dict.fromkeys(cleaned_values))


def split_display_treatments(value: str) -> list[str]:
    treatments = []
    for line in str(value or "").splitlines():
        cleaned = re.sub(r"^\s*(?:[-*]\s+|\d+[\.)]\s*)", "", line).strip()
        parts = re.split(
            r",\s+(?=(?:\d{1,2}(?:-\d{1,2})?\s+-|Kw\s+[IVX]+\s+-|Herstelling|Gebit|bijvoegen|Volledige|Behandeling ontbreekt))",
            cleaned,
        )
        treatments.extend(part.strip() for part in parts if part.strip())
    return treatments


def format_treatment_list(treatments: Iterable[str]) -> str:
    return ", ".join(unique_non_empty(treatments))


def extract_treatment_rows(file_name: str, file_bytes: bytes, planning_date: date, source_key: str) -> list[dict[str, str]]:
    fallback_name = patient_name_from_filename(file_name)
    text, pdf_error = extract_pdf_text(file_bytes)
    if pdf_error:
        return [
            {
                "_source_key": source_key,
                PATIENT_COL: fallback_name,
                TREATMENTS_COL: "",
                DATE_PERIOD_COL: format_european_date(planning_date),
                REMARKS_COL: f"{pdf_error}; Patientnaam uit bestandsnaam gehaald.",
            }
        ]

    patient_name, name_remark = guess_patient_name(text, fallback_name)
    treatments, treatment_remark = guess_treatments(text)
    remarks = "; ".join(remark for remark in [name_remark, treatment_remark] if remark)
    return [
        {
            "_source_key": source_key,
            PATIENT_COL: patient_name,
            TREATMENTS_COL: treatment,
            DATE_PERIOD_COL: format_european_date(planning_date),
            REMARKS_COL: remarks,
        }
        for treatment in treatments
    ]


def group_rows(rows: list[dict[str, str]], planning_date: date) -> list[dict[str, object]]:
    grouped: dict[str, dict[str, object]] = {}
    for row in rows:
        patient_name = str(row.get(PATIENT_COL, "") or "").strip() or "Onbekende patient"
        source_key = str(row.get("_source_key") or patient_name)
        group = grouped.setdefault(
            source_key,
            {
                PATIENT_COL: patient_name,
                TREATMENTS_COL: [],
                REMARKS_COL: [],
            },
        )
        group[TREATMENTS_COL].append(str(row.get(TREATMENTS_COL, "") or ""))
        group[REMARKS_COL].append(str(row.get(REMARKS_COL, "") or ""))

    name_counts: dict[str, int] = {}
    for group in grouped.values():
        key = simplify_for_matching(str(group[PATIENT_COL])).strip()
        if key:
            name_counts[key] = name_counts.get(key, 0) + 1

    output = []
    for number, group in enumerate(grouped.values(), start=1):
        name_key = simplify_for_matching(str(group[PATIENT_COL])).strip()
        remarks = unique_non_empty(group[REMARKS_COL])
        if name_key and name_counts.get(name_key, 0) > 1:
            remarks.append(DUPLICATE_PATIENT_NAME_REMARK)
        output.append(
            {
                PATIENT_DONE_COL: False,
                PATIENT_NR_COL: number,
                PATIENT_COL: str(group[PATIENT_COL]),
                TREATMENTS_COL: format_treatment_list(group[TREATMENTS_COL]),
                DATE_PERIOD_COL: format_european_date(planning_date),
                REMARKS_COL: "\n".join(unique_non_empty(remarks)),
            }
        )
    return output


def preview_uploaded_pdfs(files: list[tuple[str, bytes]], planning_date: date) -> list[dict[str, object]]:
    rows = []
    for index, (file_name, file_bytes) in enumerate(files, start=1):
        rows.extend(extract_treatment_rows(file_name, file_bytes, planning_date, f"{index}:{file_name}"))
    return group_rows(rows, planning_date)


def treatment_name_key(value: str) -> str:
    cleaned = re.sub(r"\s*[\[(](?:kamer|kam|k)?\s*[a-z0-9-]+[\])]\s*$", "", str(value or ""), flags=re.IGNORECASE)
    cleaned = re.sub(r"^(?:dhr|de heer|mevr|mevrouw)\.?\s+", "", cleaned, flags=re.IGNORECASE)
    tokens = re.findall(r"[a-z0-9]+", simplify_for_matching(cleaned))
    return " ".join(sorted(tokens))


def preview_treatment_pdf_files(
    files: list[tuple[str, bytes]],
    planning_date: date,
) -> list[dict[str, object]]:
    output = []
    for index, (file_name, file_bytes) in enumerate(files, start=1):
        extracted = extract_treatment_rows(file_name, file_bytes, planning_date, f"{index}:{file_name}")
        grouped = group_rows(extracted, planning_date)
        if not grouped:
            continue
        output.append(
            {
                "pdf_id": str(index),
                "file_name": file_name,
                "pdf_name": str(grouped[0].get(PATIENT_COL, "") or ""),
                "row": grouped[0],
            }
        )
    return output


def match_treatment_plans(
    roster_rows: list[dict[str, object]],
    pdf_items: list[dict[str, object]],
) -> dict[str, object]:
    roster_by_key: dict[str, list[str]] = {}
    for row in roster_rows:
        name = str(row.get(PATIENT_COL, "") or "").strip()
        key = treatment_name_key(name)
        if key:
            roster_by_key.setdefault(key, []).append(name)

    pdf_by_key: dict[str, list[dict[str, object]]] = {}
    for item in pdf_items:
        key = treatment_name_key(str(item.get("pdf_name", "") or ""))
        if key:
            pdf_by_key.setdefault(key, []).append(item)

    matched = []
    review = []
    missing_names = []
    unselected_pdfs = []

    for key, roster_names in roster_by_key.items():
        matching_pdfs = pdf_by_key.get(key, [])
        if len(roster_names) == 1 and len(matching_pdfs) == 1:
            item = matching_pdfs[0]
            item["row"][PATIENT_COL] = roster_names[0]
            matched.append(
                {
                    "pdf_id": str(item.get("pdf_id") or item["file_name"]),
                    "roster_name": roster_names[0],
                    "pdf_name": item["pdf_name"],
                    "file_name": item["file_name"],
                    "row": item["row"],
                }
            )
        elif matching_pdfs:
            reason = "Dubbele naam in de Excel-lijst." if len(roster_names) > 1 else "Meerdere PDF's met dezelfde bewonersnaam."
            for item in matching_pdfs:
                review.append(
                    {
                        "pdf_id": str(item.get("pdf_id") or item["file_name"]),
                        "roster_name": " / ".join(roster_names),
                        "pdf_name": item["pdf_name"],
                        "file_name": item["file_name"],
                        "reason": reason,
                        "row": item["row"],
                    }
                )
        else:
            missing_names.extend(roster_names)

    for key, items in pdf_by_key.items():
        if key in roster_by_key:
            continue
        for item in items:
            unselected_pdfs.append(
                {
                    "pdf_id": str(item.get("pdf_id") or item["file_name"]),
                    "pdf_name": item["pdf_name"],
                    "file_name": item["file_name"],
                    "row": item["row"],
                }
            )

    for number, item in enumerate(matched, start=1):
        item["row"][PATIENT_NR_COL] = number

    matched_names_by_pdf = {str(item["pdf_id"]): str(item["roster_name"]) for item in matched}
    review_reasons_by_pdf = {str(item["pdf_id"]): str(item["reason"]) for item in review}
    all_pdfs = [
        {
            "pdf_id": str(item.get("pdf_id") or item["file_name"]),
            "roster_name": matched_names_by_pdf.get(str(item.get("pdf_id") or item["file_name"]), ""),
            "pdf_name": item["pdf_name"],
            "file_name": item["file_name"],
            "reason": review_reasons_by_pdf.get(str(item.get("pdf_id") or item["file_name"]), ""),
            "row": item["row"],
        }
        for item in pdf_items
    ]

    return {
        "excel_count": len(roster_rows),
        "pdf_count": len(pdf_items),
        "roster_names": [str(row.get(PATIENT_COL, "") or "").strip() for row in roster_rows],
        "pdfs": all_pdfs,
        "matched": matched,
        "missing_names": missing_names,
        "unselected_pdfs": unselected_pdfs,
        "review": review,
    }


def preview_treatment_matches(
    roster_file: tuple[str, bytes],
    pdf_files: list[tuple[str, bytes]],
    planning_date: date,
) -> dict[str, object]:
    roster_rows = preview_screening_workbooks([roster_file], planning_date)
    pdf_items = preview_treatment_pdf_files(pdf_files, planning_date)
    return match_treatment_plans(roster_rows, pdf_items)


def clean_cell_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, date):
        return format_european_date(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalize_excel_header(value: object) -> str:
    normalized = simplify_for_matching(clean_cell_value(value))
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", normalized)).strip("_")


def is_screening_combined_name_header(header: str) -> bool:
    if header in {
        "patient",
        "patientnaam",
        "bewoner",
        "naam_en_voornaam",
        "naam_voornaam",
        "familienaam_voornaam",
        "voornaam_familienaam",
        "naam",
        "name",
    }:
        return True
    if header == "voornaam":
        return False
    return "voornaam" in header and any(part in header for part in ("familienaam", "achternaam", "naam"))


def is_screening_last_name_header(header: str) -> bool:
    return header in {"naam", "achternaam", "familienaam", "naam_bewoner", "naam_patient"} or (
        "voornaam" not in header and any(part in header for part in ("familienaam", "achternaam"))
    )


def is_screening_first_name_header(header: str) -> bool:
    return "voornaam" in header and not is_screening_combined_name_header(header)


def screening_header_score(values: Iterable[object]) -> int:
    headers = [normalize_excel_header(value) for value in values if clean_cell_value(value)]
    if not headers:
        return 0
    has_first_name = any(is_screening_first_name_header(header) for header in headers)
    has_last_name = any(is_screening_last_name_header(header) for header in headers)
    has_combined_name = any(is_screening_combined_name_header(header) for header in headers)
    score = 0
    if has_first_name and has_last_name:
        score += 6
    if has_combined_name:
        score += 4
    score += sum(1 for header in headers if header in {"kamer", "afdeling", "geboortedatum", "rijksregisternummer"})
    return score


def find_screening_header_row(df: pd.DataFrame) -> int | None:
    best_index = None
    best_score = 0
    for index, row in df.head(30).iterrows():
        score = screening_header_score(row.tolist())
        if score > best_score:
            best_index = int(index)
            best_score = score
    return best_index if best_score >= 3 else None


def unique_headers(values: Iterable[object]) -> list[str]:
    headers = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        header = clean_cell_value(value) or f"Kolom {index}"
        counts[header] = counts.get(header, 0) + 1
        headers.append(header if counts[header] == 1 else f"{header} {counts[header]}")
    return headers


def screening_name_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", simplify_for_matching(value)).strip()


def looks_like_date_text(value: str) -> bool:
    cleaned = str(value or "").strip()
    key = screening_name_key(cleaned)
    if re.fullmatch(r"\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?", cleaned):
        return True
    if re.fullmatch(r"\d{4}[/-]\d{1,2}[/-]\d{1,2}", cleaned):
        return True
    month_names = {
        "januari",
        "februari",
        "maart",
        "april",
        "mei",
        "juni",
        "juli",
        "augustus",
        "september",
        "oktober",
        "november",
        "december",
    }
    month_pattern = "|".join(sorted(month_names, key=len, reverse=True))
    return bool(
        re.fullmatch(rf"\d{{1,2}}\s+(?:{month_pattern})(?:\s+\d{{2,4}})?", key)
        or re.fullmatch(rf"(?:{month_pattern})\s+\d{{4}}", key)
    )


def is_screening_non_patient_name(value: str) -> bool:
    key = screening_name_key(value)
    if not key or len(key) < 3:
        return True
    if looks_like_date_text(value):
        return True
    if key in {
        "planning",
        "planning datum",
        "datum",
        "familienaam",
        "voornaam",
        "familienaam voornaam",
        "naam",
        "naam voornaam",
        "naam en voornaam",
        "bewoner",
        "bewoners",
        "patient",
        "patientnaam",
        "patient naam",
        "lijst bewoners",
        "screeninglijst",
    }:
        return True
    if key.startswith(("planning ", "datum ", "screeninglijst ")):
        return True
    return "familienaam" in key and "voornaam" in key


def find_screening_name_columns(df: pd.DataFrame) -> tuple[object | None, object | None, str]:
    headers = {column: normalize_excel_header(column) for column in df.columns}
    first_name_column = next((column for column, header in headers.items() if is_screening_first_name_header(header)), None)
    last_name_column = next(
        (
            column
            for column, header in headers.items()
            if is_screening_last_name_header(header)
        ),
        None,
    )
    if first_name_column and last_name_column and first_name_column != last_name_column:
        return last_name_column, first_name_column, ""

    for column, header in headers.items():
        if is_screening_combined_name_header(header):
            return column, None, ""

    best_column = None
    best_score = 0
    for column in df.columns:
        values = [clean_cell_value(value) for value in df[column].tolist()]
        score = sum(1 for value in values if re.search(r"[A-Za-zÀ-ÿ]", value))
        if score > best_score:
            best_column = column
            best_score = score
    if best_column is None or best_score == 0:
        return None, None, "Geen duidelijke naamkolom gevonden."
    return best_column, None, f"Naamkolom automatisch gekozen: {clean_cell_value(best_column)}."


def read_screening_table(file_name: str, file_bytes: bytes) -> pd.DataFrame:
    lower_name = file_name.lower()
    buffer = BytesIO(file_bytes)
    if lower_name.endswith(".csv"):
        raw_df = pd.read_csv(buffer, sep=None, engine="python", header=None)
    else:
        raw_df = pd.read_excel(buffer, sheet_name=0, header=None)

    header_row = find_screening_header_row(raw_df)
    if header_row is None:
        return raw_df

    df = raw_df.iloc[header_row + 1 :].copy()
    df.columns = unique_headers(raw_df.iloc[header_row].tolist())
    return df.reset_index(drop=True)


def extract_screening_rows(file_name: str, file_bytes: bytes, planning_date: date) -> list[dict[str, object]]:
    try:
        df = read_screening_table(file_name, file_bytes)
    except Exception as exc:
        raise ValueError(f"{file_name} kon niet gelezen worden als Excel/CSV-bestand.") from exc

    if df.empty:
        return []

    last_name_column, first_name_column, column_remark = find_screening_name_columns(df)
    if last_name_column is None:
        raise ValueError(f"{file_name}: {column_remark}")

    rows = []
    for _, item in df.iterrows():
        name_parts = [clean_cell_value(item.get(last_name_column))]
        if first_name_column:
            name_parts.append(clean_cell_value(item.get(first_name_column)))
        patient_name = clean_patient_name(" ".join(part for part in name_parts if part))
        if not re.search(r"[A-Za-zÀ-ÿ]", patient_name) or is_screening_non_patient_name(patient_name):
            continue
        rows.append(
            {
                PATIENT_DONE_COL: False,
                PATIENT_COL: patient_name,
                TREATMENTS_COL: SCREENING_TREATMENT,
                DATE_PERIOD_COL: format_european_date(planning_date),
                REMARKS_COL: column_remark,
            }
        )
    return rows


def preview_screening_workbooks(files: list[tuple[str, bytes]], planning_date: date) -> list[dict[str, object]]:
    rows = []
    for file_name, file_bytes in files:
        rows.extend(extract_screening_rows(file_name, file_bytes, planning_date))
    for number, row in enumerate(rows, start=1):
        row[PATIENT_NR_COL] = number
    return rows


def build_excel_file(planning_name: str, planning_date: date, location: str, rows: list[dict[str, object]], planning_type: str = "Behandeling") -> BytesIO:
    output = BytesIO()
    is_screening = simplify_for_matching(planning_type) == "screening"
    done_col = SCREENING_DONE_COL if is_screening else PATIENT_DONE_COL
    columns = [done_col, PATIENT_NR_COL, PATIENT_COL, DATE_PERIOD_COL, REMARKS_COL] if is_screening else PREVIEW_COLUMNS
    display_rows = []
    for row in rows:
        display_row = dict(row)
        display_row[done_col] = display_row.pop(PATIENT_DONE_COL, "")
        display_rows.append(display_row)
    df = pd.DataFrame(display_rows, columns=columns)
    if TREATMENTS_COL in df.columns:
        df[TREATMENTS_COL] = df[TREATMENTS_COL].apply(lambda value: "\n".join(split_display_treatments(value)))

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Planning", index=False, startrow=6)
        worksheet = writer.sheets["Planning"]
        style_excel_sheet(worksheet, planning_name, planning_date, location, len(df), done_col)
        add_checklist_validation(worksheet, len(df))

    output.seek(0)
    return output


def style_excel_sheet(worksheet, title: str, planning_date: date, location: str, row_count: int, done_col: str = PATIENT_DONE_COL) -> None:
    last_column = get_column_letter(worksheet.max_column)
    thin_border = Border(
        left=Side(style="thin", color="DDE7E2"),
        right=Side(style="thin", color="DDE7E2"),
        top=Side(style="thin", color="DDE7E2"),
        bottom=Side(style="thin", color="DDE7E2"),
    )
    worksheet.merge_cells(f"A1:{last_column}1")
    worksheet["A1"] = title or "Planning"
    worksheet["A1"].font = Font(size=16, bold=True)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet["A3"] = "Datum"
    worksheet["B3"] = format_european_date(planning_date)
    worksheet["A4"] = "Locatie"
    worksheet["B4"] = location or "Niet ingevuld"
    worksheet["A3"].font = Font(bold=True)
    worksheet["A4"].font = Font(bold=True)
    if done_col == SCREENING_DONE_COL:
        worksheet["A5"] = "Geschatte duur"
        worksheet["B5"] = f"{format_duration(row_count * 4)} op basis van 4 min/persoon"
        worksheet["A5"].font = Font(bold=True)

    header_row = 7
    header_fill = PatternFill("solid", fgColor="86C8AE")
    for cell in worksheet[header_row]:
        cell.font = Font(color="1D1D1F", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    worksheet.row_dimensions[header_row].height = 26

    widths = {
        PATIENT_DONE_COL: 16,
        SCREENING_DONE_COL: 16,
        PATIENT_NR_COL: 12,
        PATIENT_COL: 28,
        TREATMENTS_COL: 58,
        DATE_PERIOD_COL: 16,
        REMARKS_COL: 40,
    }
    for row in worksheet.iter_rows(min_row=8):
        for cell in row:
            header_value = worksheet.cell(header_row, cell.column).value
            cell.alignment = Alignment(
                horizontal="center" if header_value in {done_col, PATIENT_NR_COL} else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = thin_border
        line_count = max(len(str(cell.value).splitlines()) if cell.value is not None else 1 for cell in row)
        worksheet.row_dimensions[row[0].row].height = min(max(24, line_count * 18), 150)

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = worksheet.cell(header_row, column_cells[0].column).value
        worksheet.column_dimensions[column_letter].width = widths.get(header_value, 18)

    worksheet.freeze_panes = "A8"
    worksheet.auto_filter.ref = f"A7:{last_column}{max(7, 7 + row_count)}"


def add_checklist_validation(worksheet, row_count: int) -> None:
    if row_count == 0:
        return
    end_row = 7 + row_count
    validation = DataValidation(type="list", formula1=f'"{CHECKBOX_DONE}"', allow_blank=True)
    validation.showDropDown = False
    worksheet.add_data_validation(validation)
    validation.add(f"A8:A{end_row}")
    done_fill = PatternFill("solid", fgColor="DFF3EA")
    worksheet.conditional_formatting.add(
        f"A8:A{end_row}",
        FormulaRule(formula=[f'A8="{CHECKBOX_DONE}"'], fill=done_fill),
    )
